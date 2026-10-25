"""Contract tests for each upstream format, against trimmed copies of real payloads (tests/fixtures/).

A source that changes its layout should fail here, not show up as "n/a" or a wrong number on the page.
"""
import io
import json

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import FIXTURES
from radar import net, paths, universe
from radar.sources import constituents, external, ingest, positive_list, tracked

# ---------------- FRED ----------------

def test_fred_csv_current_format():
    s = ingest.parse_fred_csv((FIXTURES / "fred_DGS10.csv").read_text(), "DGS10")
    assert s.index[-1] == pd.Timestamp("2026-10-08") and s.iloc[-1] == 5.22
    assert pd.isna(s[pd.Timestamp("2026-09-07")])          # holiday: blank value, kept as NaN
    assert s.dtype == float and s.index.is_monotonic_increasing


def test_fred_csv_legacy_header_and_dot_for_missing():
    s = ingest.parse_fred_csv("DATE,DGS10\n2024-01-01,.\n2024-01-02,3.95\n", "DGS10")
    assert pd.isna(s.iloc[0]) and s.iloc[1] == 3.95


def test_fred_csv_without_the_series_fails():
    with pytest.raises(ValueError, match="no DGS10 column"):
        ingest.parse_fred_csv("observation_date,T10Y2Y\n2026-10-01,0.5\n", "DGS10")
    with pytest.raises(ingest.Unavailable, match=r"web page \(Service unavailable\), not CSV"):
        ingest.parse_fred_csv("<html><title>Service unavailable</title></html>", "DGS10")
    with pytest.raises(ValueError, match="no values"):
        ingest.parse_fred_csv("observation_date,DGS10\n2026-10-01,\n", "DGS10")


def test_update_fred_keeps_cached_series_when_one_fails(tmp_path, monkeypatch):
    fred_file = tmp_path / "fred.parquet"
    paths.configure(paths.Locations.flat(tmp_path))
    cached = pd.DataFrame({"T10Y2Y": [0.5]}, index=pd.to_datetime(["2026-01-02"]))
    cached.to_parquet(fred_file)

    class Resp:
        def __init__(self, text, ok=True):
            self.text, self.ok, self.headers = text, ok, {}

        def raise_for_status(self):
            if not self.ok:
                raise ConnectionError("503")

    good = (FIXTURES / "fred_DGS10.csv").read_text()
    monkeypatch.setattr(net.requests, "get",
                        lambda url, **k: Resp(good) if url.endswith("DGS10") else Resp("", ok=False))
    failed = ingest.update_fred(["DGS10", "T10Y2Y"])
    assert failed == ["T10Y2Y (ConnectionError)"]
    out = pd.read_parquet(fred_file)
    assert out["DGS10"].dropna().iloc[-1] == 5.22 and out["T10Y2Y"].dropna().iloc[0] == 0.5


# ---------------- OECD CLI ----------------

def test_oecd_cli_real_format():
    cli = external.parse_oecd_cli((FIXTURES / "oecd_cli.csv").read_text(encoding="utf-8"))
    assert set(cli.columns) == {"G20", "USA", "G4E", "CHN"}
    assert cli.index[-1] == pd.Timestamp("2026-09-01")
    assert cli.loc["2026-08-01", "USA"] == pytest.approx(100.9464)
    t = external.cli_table(cli, {"G20": "G20", "USA": "United States", "G4E": "Europe big 4", "CHN": "China"})
    assert t.iloc[0]["area"] == "G20" and t.iloc[0]["phase"] == "Slowdown"   # 100.14, falling
    assert t.set_index("area").loc["CHN", "chg_3m"] == pytest.approx(97.41035 - 98.15016)


def test_oecd_cli_refuses_two_series_per_area():
    text = (FIXTURES / "oecd_cli.csv").read_text(encoding="utf-8")
    csv = pd.read_csv(io.StringIO(text))
    other = csv[csv["REF_AREA"] == "USA"].assign(ADJUSTMENT="N", OBS_VALUE=50.0)  # e.g. a non-adjusted series
    with pytest.raises(ValueError, match="duplicate"):
        external.parse_oecd_cli(pd.concat([csv, other]).to_csv(index=False))


# ---------------- CFTC ----------------

def _cot(name):
    return json.loads((FIXTURES / name).read_text())


def test_cot_rows_real_tff_records():
    rows = external.cot_rows(_cot("cftc_tff.json"), external.COT_DATASETS["tff"][1], {"13874A": "S&P 500"})
    first = rows[rows["date"] == "2026-09-08"].set_index("group")
    assert first.loc["Asset managers", "net"] == 1153305 - 240944
    assert first.loc["Asset managers", "net_pct_oi"] == pytest.approx(44.0364, abs=1e-4)
    assert first.loc["Leveraged funds", "net_pct_oi"] == pytest.approx(-16.4639, abs=1e-4)
    assert len(rows) == 4                                       # 2 weeks x 2 groups


def test_cot_rows_real_disaggregated_records():
    rows = external.cot_rows(_cot("cftc_disagg.json"), external.COT_DATASETS["disagg"][1], {"088691": "Gold"})
    assert rows.iloc[0]["net_pct_oi"] == pytest.approx(32.8218, abs=1e-4) and rows.iloc[0]["market"] == "Gold"


def test_cot_rows_skip_futures_and_options_combined():
    raw = _cot("cftc_tff.json")
    combined = [r | {"futonly_or_combined": "Combined", "open_interest_all": "1"} for r in raw]
    rows = external.cot_rows(raw + combined, external.COT_DATASETS["tff"][1], {"13874A": "S&P 500"})
    assert len(rows) == 4 and rows["net_pct_oi"].abs().max() < 100


def test_cot_renamed_columns_fail_instead_of_storing_nothing():
    renamed = [{k.replace("asset_mgr", "assetmgr").replace("lev_money", "levmoney"): v for k, v in r.items()}
               for r in _cot("cftc_tff.json")]
    with pytest.raises(ValueError, match="CFTC format"):
        external.cot_parse(renamed, external.COT_DATASETS["tff"][1], {"13874A": "S&P 500"}, "tff")
    assert external.cot_parse([], external.COT_DATASETS["tff"][1], {}, "tff").empty   # no new weeks is fine


# ---------------- FactSet ----------------

def test_factset_report_text():
    r = external.parse_factset([(FIXTURES / "factset_excerpt.txt").read_text(encoding="utf-8")])
    assert r == {"quarter": "Q3 2026", "growth_now": 11.4, "growth_start": 9.8, "guidance_neg": 51,
                 "guidance_pos": 38, "fwd_pe": 21.3, "fwd_pe_5y": 20.1, "fwd_pe_10y": 18.7, "cy_next_growth": 13.5,
                 "rev_up_sectors": "Utilities, Industrials, Energy",
                 "rev_down_sectors": "Real Estate, Materials"}
    # the same text also holds the *trailing* P/E averages (23.9 / 22.4): the forward ones must win


def _factset_env(tmp_path, monkeypatch, text):
    paths.configure(paths.Locations.flat(tmp_path))

    class Resp:
        status_code, content = 200, b"%PDF-1.7 ..."
    monkeypatch.setattr(net.requests, "get", lambda *a, **k: Resp())
    monkeypatch.setattr(external, "factset_pages", lambda pdf: ["Key Metrics\n" + text])


def test_update_factset_stores_a_parsed_report(tmp_path, monkeypatch):
    _factset_env(tmp_path, monkeypatch, (FIXTURES / "factset_excerpt.txt").read_text(encoding="utf-8"))
    assert external.update_factset() is not None
    row = external.load_factset().iloc[-1]
    assert (row["growth_now"], row["growth_start"]) == (11.4, 9.8)


def test_update_factset_reworded_report_fails_and_is_not_stored(tmp_path, monkeypatch):
    text = (FIXTURES / "factset_excerpt.txt").read_text(encoding="utf-8").replace(
        "earnings growth rate for the S&P 500 is", "earnings growth rate for the index stands at")
    _factset_env(tmp_path, monkeypatch, text)
    with pytest.raises(ValueError, match="growth_now"):
        external.update_factset()
    assert external.load_factset().empty                       # so the next run retries the same report


# ---------------- SSGA SPY holdings ----------------

def _spy_xlsx(rows: list[list], header=("Name", "Ticker", "Identifier", "SEDOL", "Weight", "Sector",
                                         "Shares Held", "Local Currency")) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for line in (["Fund Name:", "SPDR® S&P 500® ETF Trust"], ["Ticker Symbol:", "SPY"],
                 ["Holdings:", "As of 12-Oct-2026"], []):
        ws.append(line)
    ws.append(list(header))
    for r in rows:
        ws.append(r)
    ws.append([])
    ws.append(["Past performance is not a guarantee of future results."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_spy_holdings_layout():
    rows = [["APPLE INC", "AAPL", "037833100", "2046251", 7.1, "-", 1, "USD"],
            ["BERKSHIRE HATHAWAY INC CL B", "BRK.B", "084670702", "2073390", 1.6, "-", 1, "USD"],
            ["MICROSOFT CORP", "MSFT", "594918104", "2588173", 6.5, "-", 1, "USD"],
            ["US DOLLAR", "CASH_USD", "-", "-", 0.1, "-", 1, "USD"],
            ["SOME ADR", "XYZ", "-", "-", 0.1, "-", 1, "EUR"],
            ["APPLE INC", "AAPL", "037833100", "2046251", 7.1, "-", 1, "USD"]]
    df = constituents.parse_spy_holdings(_spy_xlsx(rows), min_holdings=1)
    assert df["ticker"].tolist() == ["AAPL", "BRK-B", "MSFT"]   # Yahoo spelling, no cash, no non-USD, no dupes
    assert df.set_index("ticker").loc["AAPL", "name"] == "APPLE INC"


def test_spy_holdings_layout_change_fails():
    rows = [["APPLE INC", "AAPL", "037833100", "2046251", 7.1, "-", 1, "USD"]]
    with pytest.raises(ValueError, match="Weight"):
        constituents.parse_spy_holdings(_spy_xlsx(rows, header=("Name", "Ticker", "Identifier", "SEDOL", "Pct", "Sector",
                                                             "Shares Held", "Local Currency")))
    with pytest.raises(ValueError, match="only 1 stocks"):
        constituents.parse_spy_holdings(_spy_xlsx(rows))            # default minimum: a full index


# ---------------- SKAT positive list ----------------

POSLIST_HEADER = ["Skattemæssigt hjemsted/Tax residence", "ISIN-kode/-Code", "Navn andelsklasse/Name Shareclass",
                  "LEI-kode/-Code", "CVR/SE/TIN", "Navn afdeling/Name Sub-fund", "TIN", "Navn/Name",
                  "Registrerede år/Registered ", "Ikke registrerede år/Deregistered"]


def _poslist(tmp_data, sheets: dict[str, list[list]], header=POSLIST_HEADER):
    wb = openpyxl.Workbook()
    wb.active.title = "Forside"
    wb.active.append(["Positivliste"])
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(tmp_data / "positivliste.xlsx")


def test_positive_list_uses_the_latest_year_only(tmp_data):
    old = ["IE", "IE00BOLD0001", "", "", "", "", "", "Old fund", "2024", "2026"]
    keep = ["IE", " ie00b6r52259 ", "", "", "", "", "", "ACWI", "2018", ""]
    bad = ["IE", "not an isin", "", "", "", "", "", "Footnote", "", ""]
    _poslist(tmp_data, {"2025": [old, keep], "2026": [keep, bad]})
    isins = positive_list.isins()
    assert isins == {"IE00B6R52259"}
    assert positive_list.load().attrs["year"] == "2026"
    assert positive_list.flag("IE00BOLD0001", isins) == "no"
    assert positive_list.flag("ie00b6r52259", isins) == "yes"
    assert positive_list.flag(None, isins) == "?" and positive_list.flag("X", None) == "?"


def test_positive_list_per_income_year(tmp_data):
    """The tax estimate reads each year's own sheet: a fund that left the list is off it that year only."""
    left = ["IE", "IE00BOLD0001", "", "", "", "", "", "Old fund", "2024", "2026"]
    keep = ["IE", "IE00B6R52259", "", "", "", "", "", "ACWI", "2018", ""]
    _poslist(tmp_data, {"2025": [left, keep], "2026": [keep]})
    years = positive_list.load_years()
    assert set(years) == {2025, 2026}
    assert positive_list.status("IE00BOLD0001", 2025, years) == "yes"
    assert positive_list.status("ie00bold0001", 2026, years) == "no"
    assert positive_list.status("IE00B6R52259", 2024, years) == "no_sheet"
    assert positive_list.status(None, 2026, years) == "no_isin"
    assert positive_list.status("IE00B6R52259", 2026, None) == "no_sheet"


def test_positive_list_finds_the_isin_column_by_header(tmp_data):
    header = POSLIST_HEADER[2:3] + POSLIST_HEADER[:2] + POSLIST_HEADER[3:]   # SKAT moved a column
    _poslist(tmp_data, {"2026": [["ACWI", "IE", "IE00B6R52259", "", "", "", "", "", "", ""]]}, header)
    assert positive_list.isins() == {"IE00B6R52259"}


def test_positive_list_without_isins_fails(tmp_data):
    _poslist(tmp_data, {"2026": [["IE", "", "", "", "", "", "", "Nothing", "", ""]]})
    with pytest.raises(ValueError, match="no valid ISINs"):
        positive_list.load()


# ---------------- Yahoo download frames ----------------

def _yf_frame(n=3, start="2026-10-06"):
    idx = pd.DatetimeIndex(pd.bdate_range(start, periods=n), name="Date")
    close = np.arange(100.0, 100.0 + n)
    return pd.DataFrame({"Open": close, "High": close + 1, "Low": close - 1, "Close": close,
                         "Adj Close": close * 0.9, "Volume": 1000}, index=idx)


def test_download_handles_multi_and_single_ticker_frames(monkeypatch):
    multi = pd.concat({"AAA": _yf_frame(), "BBB": _yf_frame() * np.nan}, axis=1)   # BBB: listed but no data
    monkeypatch.setattr(ingest.yf, "download", lambda batch, **kw: multi if len(batch) > 1 else _yf_frame())
    out = ingest.download_bars(["AAA", "BBB"], pd.Timestamp("2026-10-01").date())
    assert set(out["ticker"]) == {"AAA"} and list(out.columns) == ingest.COLUMNS
    assert out["close"].tolist() == pytest.approx([90.0, 90.9, 91.8]) and out["raw_close"].iloc[0] == 100.0
    assert out["high"].iloc[0] == pytest.approx(101 * 0.9)
    single = ingest.download_bars(["AAA"], pd.Timestamp("2026-10-01").date())   # flat columns from yfinance
    assert single["ticker"].unique().tolist() == ["AAA"] and len(single) == 3


def test_download_empty_response(monkeypatch):
    monkeypatch.setattr(ingest.yf, "download", lambda batch, **kw: pd.DataFrame())
    out = ingest.download_bars(["AAA"], pd.Timestamp("2026-10-01").date())
    assert out.empty and list(out.columns) == ingest.COLUMNS


# ---------------- OECD: nothing is written before the answer is checked (hardening 1.3) ----------------

def _oecd(areas, months=24, end="2026-08"):
    rows = ["REF_AREA,TIME_PERIOD,OBS_VALUE"] + [f"{a},{m:%Y-%m},100.0" for a in areas
                                                 for m in pd.period_range(end=end, periods=months, freq="M").to_timestamp()]
    return "\n".join(rows) + "\n"


def test_oecd_header_only_answer_raises_and_writes_nothing(monkeypatch):
    from fakes import Resp, serve
    areas = list(external.cfg().get("oecd_cli", {}))
    serve(monkeypatch, {"oecd": Resp(text=_oecd(areas))})
    external.update_oecd_cli()
    before = external.CLI_FILE.read_bytes()
    serve(monkeypatch, {"oecd": Resp(text="REF_AREA,TIME_PERIOD,OBS_VALUE\n")})
    with pytest.raises(ValueError):
        external.update_oecd_cli()
    assert external.CLI_FILE.read_bytes() == before


def test_oecd_area_missing_from_the_answer_keeps_its_stored_values(monkeypatch):
    from fakes import Resp, serve
    areas = list(external.cfg().get("oecd_cli", {}))
    serve(monkeypatch, {"oecd": Resp(text=_oecd(areas))})
    external.update_oecd_cli()
    serve(monkeypatch, {"oecd": Resp(text=_oecd(areas[1:], end="2026-09"))})
    msg = external.update_oecd_cli()
    df = pd.read_parquet(external.CLI_FILE)
    assert areas[0] in df.columns and df[areas[0]].notna().sum() == 24
    assert f"kept the stored {areas[0]}" in msg


# ---------------- reference caches keep answers, not failures (hardening 1.7) ----------------

def test_a_failed_name_lookup_is_not_cached(tmp_data, monkeypatch):
    from fakes import fake_yf
    monkeypatch.setattr(tracked, "etfs", lambda: [])
    yf = fake_yf(monkeypatch, info={"A": RuntimeError("429"), "B": {}, "C": {"longName": "Cee Corp"}})
    assert universe.names(["A", "B", "C"]) == {"A": "A", "B": "B", "C": "Cee Corp"}   # shown as the ticker...
    cached = pd.read_csv(paths.CACHE / "names.csv")
    assert list(cached["ticker"]) == ["C"]                                             # ...but not remembered
    yf.asked.clear()
    universe.names(["A", "B", "C"])
    assert sorted(t for _, t in yf.asked) == ["A", "B"]                                # so asked again


def test_a_sector_is_cached_only_from_a_real_answer(tmp_data, monkeypatch):
    from fakes import fake_yf
    fake_yf(monkeypatch, info={"S": {"quoteType": "EQUITY", "sector": "Technology", "industry": "Chips"},
                               "E": {"quoteType": "ETF"}, "T": {}})
    universe.sectors(["S", "E", "T"])
    cached = pd.read_csv(paths.CACHE / "sectors.csv", keep_default_na=False)
    assert sorted(cached["ticker"]) == ["E", "S"]          # T's empty answer (throttled?) is asked again later


def test_acwi_region_keys_match():
    """ACWI's holdings and the radar's own stocks are mapped to the same regions, so an active weight compares like
    with like: every region one side can produce, the other can too."""
    df = constituents.parse_acwi_holdings((FIXTURES / "ishares_ssac_holdings.csv").read_text(encoding="utf-8"),
                                          min_rows=10)
    countries = set(df["country"]) | {"Vereinigte Staaten", "Dänemark", "Japan", "Kanada", "Australien", ""}
    acwi = {constituents.acwi_region(c) for c in countries}
    assert acwi == set(universe.REGIONS) and set(df["region"]) <= set(universe.REGIONS)
    assert set(tracked.stock_groups(offline=True)) <= set(universe.REGIONS)
    # each suffix region and both fallbacks are reachable from region_for
    samples = {"AAPL": "us", "SHOP.TO": "other", "BHP.AX": "other", "PETR4.SA": "other", "NOKIA.HE": "nordic",
               "XX.PR": "europe", "XX.JK": "asia", "9984.T": "asia"}
    assert {t: universe.region_for(t) for t in samples} == samples
    assert {universe.region_for(t) for t in samples} == set(universe.REGIONS)


# ---------------- ACWI's weights over time (attribution) ----------------

def _ssac(asof_line='Fondsposition per,"14.Okt.2026"'):
    text = (FIXTURES / "ishares_ssac_holdings.csv").read_text(encoding="utf-8")
    return constituents.parse_acwi_holdings(text.replace('Fondsposition per,"14.Okt.2026"', asof_line), min_rows=10)


def test_acwi_asof_is_parsed_from_the_file():
    assert set(_ssac()["asof"]) == {"2026-10-14"}
    assert set(_ssac('Fondsposition per,"03.Mär.2025"')["asof"]) == {"2025-03-03"}
    assert set(_ssac("")["asof"]) == {None}


def test_acwi_segments_add_up_per_dimension():
    seg = constituents.acwi_segments(_ssac())
    for dim in ("sector", "region", "country", "currency", "sector_currency", "region_currency", "region_country"):
        assert seg.loc[seg["dim"] == dim, "weight"].sum() == pytest.approx(1.0)
    assert "United States" in set(seg.loc[seg["dim"] == "country", "key"])
    assert "us|United States" in set(seg.loc[seg["dim"] == "region_country", "key"])


def test_acwi_segment_history_is_append_only_once_a_week(tmp_data):
    first = constituents.record_acwi_segments(_ssac('Fondsposition per,"14.Okt.2026"'))
    assert "recorded" in first
    before = constituents.acwi_segment_history()
    assert "already" in constituents.record_acwi_segments(_ssac('Fondsposition per,"16.Okt.2026"'))  # same week
    assert "older" in constituents.record_acwi_segments(_ssac('Fondsposition per,"02.Okt.2026"'))
    pd.testing.assert_frame_equal(constituents.acwi_segment_history(), before)
    constituents.record_acwi_segments(_ssac('Fondsposition per,"21.Okt.2026"'))
    after = constituents.acwi_segment_history()
    assert sorted(after["asof"].unique()) == ["2026-10-14", "2026-10-21"]
    pd.testing.assert_frame_equal(after.iloc[: len(before)].reset_index(drop=True), before)   # earlier rows unchanged
    with pytest.raises(ValueError, match="no date"):
        constituents.record_acwi_segments(_ssac(""))


def test_acwi_keys_are_ones_the_attribution_can_use():
    """Every sector ACWI's file can name has a proxy and is a sector the radar's classification uses; every country
    the region map knows has an English name (the country proxies are keyed by it)."""
    from radar.config import load_yaml
    from radar.discovery.sources import ISHARES_SECTORS

    conf = load_yaml("universe.yaml")
    assert set(ISHARES_SECTORS.values()) == set(conf["attribution"]["sectors"])
    assert set(ISHARES_SECTORS.values()) <= set(conf["sector_etfs"]["us"])
    seen = ["Vereinigte Staaten", "Japan", "Taiwan", "Vereinigtes Königreich", "Kanada", "Deutschland", "Korea",
            "Schweiz", "Irland", "Frankreich", "China", "Australien", "Niederlande", "Spanien", "Schweden", "Italien",
            "Dänemark", "Singapur", "Südafrika", "Hongkong", "Belgien", "Finnland", "Israel", "Mexiko", "Norwegen",
            "Thailand", "Ver. Arabische Emirate", "Malaysia", "Polen", "Österreich", "Indonesien", "Portugal",
            "Griechenland", "Peru", "Kuwait", "Ungarn", "Qatar", "Neuseeland", "Chile", "Türkei", "Philippinen",
            "Kolumbien", "Indien", "Tschechien", "Ägypten", "Russland"]   # every country in the October 2026 file
    assert all(c in constituents.COUNTRY_EN for c in seen)
    regions = conf["attribution"]["regions"]
    assert set(regions) == set(universe.REGIONS)
    for r, countries in regions.items():
        assert all(constituents.acwi_region(de) == r for de, en in constituents.COUNTRY_EN.items() if en in countries)


def test_the_update_step_records_the_cached_files_week(tmp_data, monkeypatch):
    from radar.config import write_csv

    def offline(*a, **k):
        raise ConnectionError("offline")
    monkeypatch.setattr(net, "get", offline)
    with pytest.raises(RuntimeError, match="no ACWI holdings"):
        constituents.update_acwi_segments()
    write_csv(_ssac(), constituents.ACWI_FILE, index=False)        # fresh: read from the cache, no request
    assert "2026-10-14" in constituents.update_acwi_segments()
    assert "already" in constituents.update_acwi_segments()
    write_csv(_ssac().drop(columns="asof"), constituents.ACWI_FILE, index=False)   # a cache from before the date
    with pytest.raises(ValueError, match="no date"):
        constituents.update_acwi_segments()
